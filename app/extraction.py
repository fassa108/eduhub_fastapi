import io
import multiprocessing
from concurrent.futures import ProcessPoolExecutor
from concurrent.futures import TimeoutError as DelaiDepasse

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.config import settings


class ExtractionError(Exception):
    """Erreur de base pour tout problème d'extraction."""


class FormatNonSupporte(ExtractionError):
    pass


class FichierIllisible(ExtractionError):
    pass


class TexteVide(ExtractionError):
    pass


class TropDePages(ExtractionError):
    pass


class ExtractionTropLongue(ExtractionError):
    pass


def _verifier_nb_pages(nb: int, unite: str) -> None:
    if nb > settings.nb_pages_max:
        raise TropDePages(f"{nb} {unite}, maximum {settings.nb_pages_max}")


def extraire_txt(contenu: bytes) -> str:
    try:
        return contenu.decode("utf-8")
    except UnicodeDecodeError as e:
        raise FichierIllisible("Le fichier texte n'est pas en UTF-8") from e


def extraire_pdf(contenu: bytes) -> str:
    try:
        document = pymupdf.open(stream=contenu, filetype="pdf")
    except Exception as e:
        raise FichierIllisible("PDF illisible ou corrompu") from e
    with document:
        _verifier_nb_pages(document.page_count, "pages")
        try:
            pages = [page.get_text() for page in document]
        except Exception as e:
            raise FichierIllisible("PDF illisible ou corrompu") from e
    return "\n".join(pages)


def extraire_docx(contenu: bytes) -> str:
    try:
        document = Document(io.BytesIO(contenu))
    except Exception as e:
        raise FichierIllisible("DOCX illisible ou corrompu") from e

    lignes = [p.text for p in document.paragraphs]
    for tableau in document.tables:
        for ligne in tableau.rows:
            lignes.append(" | ".join(cellule.text for cellule in ligne.cells))
    return "\n".join(lignes)


def _textes_des_formes(formes) -> list[str]:
    """Texte des zones, des tableaux et des groupes de formes (récursif)."""
    lignes = []
    for forme in formes:
        if forme.shape_type == MSO_SHAPE_TYPE.GROUP:
            lignes.extend(_textes_des_formes(forme.shapes))
        elif forme.has_text_frame:
            lignes.append(forme.text_frame.text)
        elif forme.has_table:
            for ligne in forme.table.rows:
                lignes.append(" | ".join(cellule.text for cellule in ligne.cells))
    return lignes


def extraire_pptx(contenu: bytes) -> str:
    try:
        presentation = Presentation(io.BytesIO(contenu))
    except Exception as e:
        raise FichierIllisible("PPTX illisible ou corrompu") from e

    _verifier_nb_pages(len(presentation.slides), "diapositives")
    lignes = []
    for numero, diapo in enumerate(presentation.slides, start=1):
        textes = _textes_des_formes(diapo.shapes)
        if diapo.has_notes_slide:
            textes.append(diapo.notes_slide.notes_text_frame.text)
        textes = [t for t in textes if t.strip()]
        if textes:
            lignes.append(f"--- Diapositive {numero} ---")
            lignes.extend(textes)
    # python-pptx rend les sauts de ligne manuels par \v
    return "\n".join(lignes).replace("\v", "\n")


EXTRACTEURS = {
    "txt": extraire_txt,
    "pdf": extraire_pdf,
    "docx": extraire_docx,
    "pptx": extraire_pptx,
}


def extraire_texte(nom: str, contenu: bytes) -> str:
    extension = nom.lower().rsplit(".", 1)[-1] if "." in nom else ""
    extracteur = EXTRACTEURS.get(extension)
    if extracteur is None:
        raise FormatNonSupporte(f"Format non supporté : '{extension}'")

    texte = extracteur(contenu).strip()
    if not texte:
        raise TexteVide(f"Aucun texte extractible dans '{nom}'")
    return texte


def extraire_texte_avec_delai(nom: str, contenu: bytes, extracteur=extraire_texte) -> str:
    """
    Extrait dans un processus à part, arrêté net au-delà du délai : un
    fichier piégé ne peut pas bloquer le service.
    """
    contexte = multiprocessing.get_context("forkserver")
    with ProcessPoolExecutor(max_workers=1, mp_context=contexte) as pool:
        futur = pool.submit(extracteur, nom, contenu)
        try:
            return futur.result(timeout=settings.extraction_delai_s)
        except DelaiDepasse:
            pool.kill_workers()
            raise ExtractionTropLongue(
                f"Extraction de '{nom}' trop longue (plus de {settings.extraction_delai_s} s)"
            ) from None

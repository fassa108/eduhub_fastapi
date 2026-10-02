import io

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.extraction import FormatNonSupporte, TexteVide, extraire_texte


def test_txt():
    assert extraire_texte("cours.txt", "Bonjour le monde".encode()) == "Bonjour le monde"


def test_format_non_supporte():
    with pytest.raises(FormatNonSupporte):
        extraire_texte("cours.odp", b"contenu")


def test_texte_vide():
    with pytest.raises(TexteVide):
        extraire_texte("vide.txt", b"   ")


def test_docx():
    document = Document()
    document.add_paragraph("Introduction à Python")
    tampon = io.BytesIO()
    document.save(tampon)

    assert "Introduction à Python" in extraire_texte("cours.docx", tampon.getvalue())


def test_pdf():
    chemin = "tests/fixtures/UML - Diagramme de classe.pdf"
    with open(chemin, "rb") as fichier:
        contenu = fichier.read()

    texte = extraire_texte(chemin, contenu)
    texte_normalise = " ".join(texte.split())

    assert "Diagramme de classe" in texte_normalise


def _pptx(remplir) -> bytes:
    presentation = Presentation()
    remplir(presentation)
    tampon = io.BytesIO()
    presentation.save(tampon)
    return tampon.getvalue()


def test_pptx_texte_tableau_groupe_et_notes():
    def remplir(presentation):
        diapo = presentation.slides.add_slide(presentation.slide_layouts[5])
        diapo.shapes.title.text = "Les design patterns"
        tableau = diapo.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        tableau.cell(0, 0).text = "Pattern"
        tableau.cell(0, 1).text = "Usage"
        groupe = diapo.shapes.add_group_shape()
        zone = groupe.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1))
        zone.text_frame.text = "Singleton"
        diapo.notes_slide.notes_text_frame.text = "Insister sur l'injection de dépendances"

    texte = extraire_texte("cours.pptx", _pptx(remplir))

    assert "--- Diapositive 1 ---" in texte
    assert "Les design patterns" in texte
    assert "Pattern | Usage" in texte
    assert "Singleton" in texte
    assert "Insister sur l'injection de dépendances" in texte


def test_pptx_sans_texte():
    def remplir(presentation):
        presentation.slides.add_slide(presentation.slide_layouts[6])

    with pytest.raises(TexteVide):
        extraire_texte("vide.pptx", _pptx(remplir))
